"""Unit tests for scanner.workbook — Excel report building."""
from __future__ import annotations

import tempfile
import unittest
from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook

from scanner.workbook import merge, summary_rows, write_workbook


class TestMerge(unittest.TestCase):
    """Test merge function."""

    def test_merge_keeps_all_stocks(self):
        """Left join preserves all stocks from enriched."""
        enriched = pd.DataFrame({
            "symbol": ["STOCK1", "STOCK2", "STOCK3"],
            "company": ["Co1", "Co2", "Co3"],
            "price": [100, 200, 300],
        })
        conclusions = pd.DataFrame({
            "symbol": ["STOCK1", "STOCK3"],
            "stance": ["CONSTRUCTIVE", "CAUTIOUS"],
        })

        result = merge(enriched, conclusions)
        self.assertEqual(len(result), 3)
        self.assertIn("STOCK2", result["symbol"].values)
        self.assertTrue(result.loc[result["symbol"] == "STOCK2", "stance"].isna().all())

    def test_merge_keeps_all_columns(self):
        """Merge preserves all columns from both DataFrames."""
        enriched = pd.DataFrame({
            "symbol": ["S1"],
            "col_a": [1],
            "col_b": [2],
            "col_c": [3],
        })
        conclusions = pd.DataFrame({
            "symbol": ["S1"],
            "col_d": ["text"],
            "col_e": [4.5],
        })

        result = merge(enriched, conclusions)
        expected_cols = {"symbol", "col_a", "col_b", "col_c", "col_d", "col_e"}
        self.assertEqual(set(result.columns), expected_cols)

    def test_merge_empty_enriched(self):
        """Merge with empty enriched returns empty."""
        enriched = pd.DataFrame()
        conclusions = pd.DataFrame({
            "symbol": ["S1"],
            "stance": ["MIXED"],
        })

        result = merge(enriched, conclusions)
        self.assertEqual(len(result), 0)

    def test_merge_empty_conclusions(self):
        """Merge with empty conclusions adds empty conclusion columns."""
        enriched = pd.DataFrame({
            "symbol": ["S1", "S2"],
            "company": ["Co1", "Co2"],
        })
        conclusions = pd.DataFrame()

        result = merge(enriched, conclusions)
        self.assertEqual(len(result), 2)
        self.assertIn("stance", result.columns)
        self.assertIn("conclusion", result.columns)
        self.assertTrue(result["stance"].isna().all())


class TestSummaryRows(unittest.TestCase):
    """Test summary_rows function."""

    def test_header_info(self):
        """Summary includes as_of date and generated timestamp."""
        df = pd.DataFrame({
            "symbol": ["S1"],
            "bucket": ["0-5%"],
            "company": ["Co1"],
        })
        rows = summary_rows(df, "2026-09-29")

        # Find header rows
        row_dict = {r[0]: r[1] for r in rows}
        self.assertEqual(row_dict["As of"], "2026-09-29")
        self.assertIn("Generated", row_dict)
        self.assertEqual(row_dict["Total stocks"], 1)

    def test_bucket_counts(self):
        """Summary includes bucket counts."""
        df = pd.DataFrame({
            "symbol": ["S1", "S2", "S3"],
            "bucket": ["NEW_ATH_TODAY", "NEW_ATH_TODAY", "0-5%"],
            "company": ["C1", "C2", "C3"],
        })
        rows = summary_rows(df, "2026-09-29")

        row_dict = {r[0]: r[1] for r in rows}
        self.assertEqual(row_dict.get("  NEW_ATH_TODAY"), 2)
        self.assertEqual(row_dict.get("  0-5%"), 1)

    def test_stance_counts(self):
        """Summary includes stance counts."""
        df = pd.DataFrame({
            "symbol": ["S1", "S2", "S3"],
            "stance": ["CONSTRUCTIVE", "MIXED", "CONSTRUCTIVE"],
            "bucket": ["0-5%", "0-5%", "0-5%"],
            "company": ["C1", "C2", "C3"],
            "conclusion": ["text", "text", "text"],
        })
        rows = summary_rows(df, "2026-09-29")

        row_dict = {r[0]: r[1] for r in rows}
        self.assertEqual(row_dict.get("  CONSTRUCTIVE"), 2)
        self.assertEqual(row_dict.get("  MIXED"), 1)

    def test_industry_counts(self):
        """Summary includes top industries."""
        df = pd.DataFrame({
            "symbol": ["S1", "S2", "S3"],
            "industry": ["Tech", "Tech", "Auto"],
            "bucket": ["0-5%", "0-5%", "0-5%"],
            "company": ["C1", "C2", "C3"],
        })
        rows = summary_rows(df, "2026-09-29")

        row_dict = {r[0]: r[1] for r in rows}
        self.assertEqual(row_dict.get("  Tech"), 2)
        self.assertEqual(row_dict.get("  Auto"), 1)

    def test_signal_counts(self):
        """Summary includes signal counts."""
        df = pd.DataFrame({
            "symbol": ["S1", "S2"],
            "orderbook_signal": [True, False],
            "expansion_signal": [True, True],
            "recent_ipo": [False, False],
            "above_ema50": [True, True],
            "above_sma200": [False, True],
            "bucket": ["0-5%", "0-5%"],
            "company": ["C1", "C2"],
        })
        rows = summary_rows(df, "2026-09-29")

        row_dict = {r[0]: r[1] for r in rows}
        self.assertEqual(row_dict.get("  Orderbook signal"), 1)
        self.assertEqual(row_dict.get("  Expansion signal"), 2)
        self.assertEqual(row_dict.get("  Above 50 EMA"), 2)

    def test_empty_dataframe(self):
        """Summary handles empty DataFrame gracefully."""
        df = pd.DataFrame()
        rows = summary_rows(df, "2026-09-29")

        row_dict = {r[0]: r[1] for r in rows}
        self.assertEqual(row_dict["Total stocks"], 0)


class TestWriteWorkbook(unittest.TestCase):
    """Test write_workbook function."""

    def setUp(self):
        """Create temp directory for test output."""
        self.temp_dir = tempfile.mkdtemp()
        self.temp_path = Path(self.temp_dir)

    def test_sheet_names_and_order(self):
        """Workbook has Stocks and Summary sheets in that order."""
        enriched = pd.DataFrame({
            "symbol": ["S1"],
            "company": ["Co1"],
            "bucket": ["0-5%"],
            "distance_pct": [2.5],
        })
        conclusions = pd.DataFrame()

        output_path = write_workbook(enriched, conclusions, "2026-09-29", str(self.temp_path))

        wb = load_workbook(output_path)
        sheet_names = wb.sheetnames
        self.assertEqual(sheet_names[0], "Stocks")
        self.assertEqual(sheet_names[1], "Summary")
        self.assertEqual(len(sheet_names), 2)

    def test_stocks_sheet_has_correct_data(self):
        """Stocks sheet contains the merged and sorted data."""
        enriched = pd.DataFrame({
            "symbol": ["S1", "S2"],
            "company": ["Co1", "Co2"],
            "bucket": ["0-5%", "NEW_ATH_TODAY"],
            "distance_pct": [2.5, 0.5],
            "industry": ["Tech", "Auto"],
        })
        conclusions = pd.DataFrame()

        output_path = write_workbook(enriched, conclusions, "2026-09-29", str(self.temp_path))

        df = pd.read_excel(output_path, sheet_name="Stocks")
        # Should be sorted: NEW_ATH_TODAY first, then 0-5%
        self.assertEqual(df.iloc[0]["bucket"], "NEW_ATH_TODAY")
        self.assertEqual(df.iloc[1]["bucket"], "0-5%")

    def test_summary_sheet_exists(self):
        """Summary sheet is created with rows."""
        enriched = pd.DataFrame({
            "symbol": ["S1"],
            "company": ["Co1"],
            "bucket": ["0-5%"],
            "distance_pct": [2.5],
        })
        conclusions = pd.DataFrame()

        output_path = write_workbook(enriched, conclusions, "2026-09-29", str(self.temp_path))

        df = pd.read_excel(output_path, sheet_name="Summary")
        self.assertGreater(len(df), 0)
        self.assertIn("Metric", df.columns)
        self.assertIn("Value", df.columns)

    def test_freeze_panes_set(self):
        """Stocks sheet has freeze panes set to C2."""
        enriched = pd.DataFrame({
            "symbol": ["S1"],
            "company": ["Co1"],
            "bucket": ["0-5%"],
            "distance_pct": [2.5],
        })
        conclusions = pd.DataFrame()

        output_path = write_workbook(enriched, conclusions, "2026-09-29", str(self.temp_path))

        wb = load_workbook(output_path)
        stocks_sheet = wb["Stocks"]
        self.assertEqual(stocks_sheet.freeze_panes, "C2")

    def test_autofilter_set(self):
        """Stocks sheet has autofilter over data range."""
        enriched = pd.DataFrame({
            "symbol": ["S1", "S2"],
            "company": ["Co1", "Co2"],
            "bucket": ["0-5%", "5-10%"],
            "distance_pct": [2.5, 7.5],
        })
        conclusions = pd.DataFrame()

        output_path = write_workbook(enriched, conclusions, "2026-09-29", str(self.temp_path))

        wb = load_workbook(output_path)
        stocks_sheet = wb["Stocks"]
        # Autofilter should be set on row 1 covering all columns
        self.assertIsNotNone(stocks_sheet.auto_filter.ref)
        self.assertIn("A1", stocks_sheet.auto_filter.ref)

    def test_returns_path(self):
        """write_workbook returns the Path object."""
        enriched = pd.DataFrame({
            "symbol": ["S1"],
            "company": ["Co1"],
            "bucket": ["0-5%"],
            "distance_pct": [2.5],
        })
        conclusions = pd.DataFrame()

        result = write_workbook(enriched, conclusions, "2026-09-29", str(self.temp_path))

        self.assertIsInstance(result, Path)
        self.assertTrue(result.exists())
        self.assertTrue(str(result).endswith(".xlsx"))

    def test_with_conclusions(self):
        """Merge includes conclusion columns when provided."""
        enriched = pd.DataFrame({
            "symbol": ["S1", "S2"],
            "company": ["Co1", "Co2"],
            "bucket": ["0-5%", "0-5%"],
            "distance_pct": [2.5, 7.5],
        })
        conclusions = pd.DataFrame({
            "symbol": ["S1"],
            "stance": ["CONSTRUCTIVE"],
            "conclusion": ["Good stock"],
        })

        output_path = write_workbook(enriched, conclusions, "2026-09-29", str(self.temp_path))

        df = pd.read_excel(output_path, sheet_name="Stocks")
        self.assertIn("stance", df.columns)
        self.assertIn("conclusion", df.columns)
        self.assertEqual(df.iloc[0]["stance"], "CONSTRUCTIVE")
        # S2 has no conclusion, should be NaN
        s2_row = df[df["symbol"] == "S2"].iloc[0]
        self.assertTrue(pd.isna(s2_row["stance"]))

    def test_missing_optional_columns(self):
        """Handles missing trend/factor columns gracefully."""
        enriched = pd.DataFrame({
            "symbol": ["S1"],
            "company": ["Co1"],
            "bucket": ["0-5%"],
            "distance_pct": [2.5],
            # Missing trend, factor columns
        })
        conclusions = pd.DataFrame()

        # Should not raise
        output_path = write_workbook(enriched, conclusions, "2026-09-29", str(self.temp_path))
        self.assertTrue(output_path.exists())


if __name__ == "__main__":
    unittest.main()
