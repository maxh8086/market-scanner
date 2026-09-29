"""Excel workbook builder for final ATH scan reports."""
from __future__ import annotations

from datetime import datetime
from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Font
from openpyxl.utils import get_column_letter

from scanner.ath import bucket_labels


def _bucket_order(max_dist: float = 26.0) -> list[str]:
    """Return bucket order for sorting."""
    return ["NEW_ATH_TODAY"] + bucket_labels(max_dist)


def merge(enriched_df: pd.DataFrame, conclusions_df: pd.DataFrame) -> pd.DataFrame:
    """Merge enriched scan data with conclusions via left join on symbol.

    Preserves all rows from enriched_df, adds conclusion columns (which will be
    empty for stocks without conclusions).
    """
    if enriched_df.empty:
        return enriched_df.copy()

    result = enriched_df.copy()
    if not conclusions_df.empty:
        result = result.merge(conclusions_df, on="symbol", how="left")
    else:
        # Add conclusion columns as empty
        for col in ["stance", "conclusion", "positives", "concerns", "data_gaps",
                    "ungrounded_points", "unverified_numbers", "speculative_phrases",
                    "stance_score", "stance_rules", "model"]:
            result[col] = None

    return result


def _sort_stocks(df: pd.DataFrame, max_dist: float = 26.0) -> pd.DataFrame:
    """Sort by bucket order, then by distance_pct ascending."""
    if df.empty:
        return df

    order = _bucket_order(max_dist)
    df = df.copy()
    df["_bucket_sort"] = df["bucket"].map(
        lambda b: order.index(b) if b in order else len(order)
    )
    df = df.sort_values(["_bucket_sort", "distance_pct"])
    df = df.drop(columns="_bucket_sort")
    return df.reset_index(drop=True)


def _organize_columns(df: pd.DataFrame) -> list[str]:
    """Return column order: identity, scan, trend/factors, then conclusion cols, then rest, then QA."""
    identity_cols = ["symbol", "company"]
    scan_cols = [
        "industry", "ltp", "ath_high", "ath_date", "distance_pct", "bucket",
        "days_since_ath", "listed_days", "recent_ipo"
    ]
    trend_cols = ["ema50", "sma200", "above_ema50", "above_sma200", "trend"]
    factor_cols = ["bonus_issue", "demerger", "orderbook_signal", "expansion_signal", "factors_matched"]
    conclusion_cols = ["stance", "stance_score", "conclusion", "positives", "concerns", "data_gaps"]
    qa_cols = ["ungrounded_points", "unverified_numbers", "speculative_phrases", "stance_rules", "model"]

    result = identity_cols.copy()
    result.extend(c for c in scan_cols if c in df.columns)
    result.extend(c for c in trend_cols if c in df.columns)
    result.extend(c for c in factor_cols if c in df.columns)
    result.extend(c for c in conclusion_cols if c in df.columns)

    # Add remaining enrichment columns (not yet listed)
    seen = set(result)
    for col in df.columns:
        if col not in seen and col not in qa_cols:
            result.append(col)

    # Add QA columns at end
    result.extend(c for c in qa_cols if c in df.columns)

    return result


def summary_rows(df: pd.DataFrame, as_of: str) -> list[tuple]:
    """Build summary sheet rows: metrics and counts."""
    rows = []
    now = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    # Header
    rows.append(("ATH Scan Report", ""))
    rows.append(("As of", as_of))
    rows.append(("Generated", now))
    rows.append(("Total stocks", len(df)))
    with_conclusion = df["conclusion"].notna().sum() if "conclusion" in df.columns else 0
    rows.append(("Stocks with conclusions", int(with_conclusion)))
    rows.append(("", ""))

    # By bucket
    rows.append(("By Bucket", ""))
    order = _bucket_order()
    for bucket in order:
        count = int((df["bucket"] == bucket).sum()) if not df.empty else 0
        rows.append((f"  {bucket}", count))
    rows.append(("", ""))

    # By stance
    if "stance" in df.columns:
        rows.append(("By Stance", ""))
        for stance in ["CONSTRUCTIVE", "MIXED", "CAUTIOUS"]:
            count = int((df["stance"] == stance).sum())
            if count > 0:
                rows.append((f"  {stance}", count))
        no_conclusion = df["stance"].isna().sum()
        if no_conclusion > 0:
            rows.append(("  (no conclusion)", int(no_conclusion)))
        rows.append(("", ""))

    # Bucket x Stance cross-tab
    if "stance" in df.columns and not df.empty:
        rows.append(("Bucket × Stance", "CONSTRUCTIVE", "MIXED", "CAUTIOUS"))
        stances = ["CONSTRUCTIVE", "MIXED", "CAUTIOUS"]
        for bucket in order:
            bucket_df = df[df["bucket"] == bucket]
            if not bucket_df.empty:
                rows.append((f"  {bucket}",
                             *(int((bucket_df["stance"] == s).sum()) for s in stances)))
        rows.append(("", ""))

    # By trend
    if "trend" in df.columns and not df.empty:
        rows.append(("By Trend (50 EMA / 200 SMA)", ""))
        for label, count in df["trend"].value_counts().items():
            rows.append((f"  {label}", int(count)))
        rows.append(("", ""))

    # Top 15 industries
    if not df.empty and "industry" in df.columns:
        rows.append(("Top 15 Industries", ""))
        for industry, count in df["industry"].value_counts().head(15).items():
            rows.append((f"  {industry}", int(count)))
        rows.append(("", ""))

    # Signal counts
    rows.append(("Signals & Characteristics", ""))
    signals = [
        ("orderbook_signal", "Orderbook signal"),
        ("expansion_signal", "Expansion signal"),
        ("recent_ipo", "Recent IPO"),
        ("above_ema50", "Above 50 EMA"),
        ("above_sma200", "Above 200 SMA"),
    ]
    for col, label in signals:
        if col in df.columns and not df.empty:
            count = df[col].fillna(0).astype(bool).sum()
            rows.append((f"  {label}", int(count)))

    # NEW_ATH_TODAY bucket is a signal too
    if not df.empty:
        new_ath_today = int((df["bucket"] == "NEW_ATH_TODAY").sum())
        if new_ath_today > 0:
            rows.append(("  New ATH today", new_ath_today))

    rows.append(("", ""))

    # Vol regime if present
    if "vol_regime" in df.columns and not df.empty:
        rows.append(("Volatility Regime", ""))
        for regime, count in df["vol_regime"].value_counts().items():
            rows.append((f"  {regime}", int(count)))
        rows.append(("", ""))

    # Data quality
    if not df.empty:
        rows.append(("Data Quality (QA)", ""))
        ungrounded_count = 0
        unverified_count = 0
        speculative_count = 0

        if "ungrounded_points" in df.columns:
            ungrounded_count = int((df["ungrounded_points"] > 0).sum())
        if "unverified_numbers" in df.columns:
            unverified_count = int((df["unverified_numbers"].fillna("").astype(str).str.len() > 0).sum())
        if "speculative_phrases" in df.columns:
            speculative_count = int((df["speculative_phrases"].fillna("").astype(str).str.len() > 0).sum())

        if ungrounded_count > 0:
            rows.append((f"  Conclusions with ungrounded points", ungrounded_count))
        if unverified_count > 0:
            rows.append((f"  Conclusions with unverified numbers", unverified_count))
        if speculative_count > 0:
            rows.append((f"  Conclusions with speculative phrases", speculative_count))

        rows.append(("", ""))

    # Closing note
    rows.append(("Note", ""))
    rows.append(("", "Conclusions are a reading aid over stored data, not investment advice;"))
    rows.append(("", "ranking is by distance from all-time high only."))

    return rows


def write_workbook(
    enriched_df: pd.DataFrame,
    conclusions_df: pd.DataFrame,
    as_of: str,
    output_dir: str = "output"
) -> Path:
    """Write Excel workbook with Stocks and Summary sheets.

    Handles locked files by writing to a timestamped filename instead.
    Returns the path written.
    """
    out_dir = Path(output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)

    # Merge data
    merged_df = merge(enriched_df, conclusions_df)
    merged_df = _sort_stocks(merged_df)

    # Organize column order
    cols_ordered = _organize_columns(merged_df)
    # Keep only columns that exist
    cols_ordered = [c for c in cols_ordered if c in merged_df.columns]
    merged_df = merged_df[cols_ordered]

    # Prepare summary
    summary_data = summary_rows(merged_df, as_of)
    width = max(len(r) for r in summary_data)
    padded = [tuple(r) + ("",) * (width - len(r)) for r in summary_data]
    summary_df = pd.DataFrame(padded, columns=["Metric", "Value", " ", "  "][:width])

    # Try the target path, fall back to timestamped if locked
    xlsx_path = out_dir / f"ath_report_{as_of}.xlsx"
    try:
        with open(xlsx_path, "a+b"):
            pass
    except PermissionError:
        stamp = datetime.now().strftime("%H%M%S")
        xlsx_path = out_dir / f"ath_report_{as_of}_{stamp}.xlsx"

    # Write workbook
    with pd.ExcelWriter(xlsx_path, engine="openpyxl") as writer:
        merged_df.to_excel(writer, sheet_name="Stocks", index=False)
        summary_df.to_excel(writer, sheet_name="Summary", index=False)

    # Format workbook: freeze panes, autofilter, widths, bold headers
    wb = load_workbook(xlsx_path)

    # Stocks sheet
    stocks_sheet = wb["Stocks"]

    # Freeze panes: row 1 (headers), first 2 columns (symbol, company)
    stocks_sheet.freeze_panes = "C2"

    # Autofilter over entire data range
    if not merged_df.empty:
        max_row = len(merged_df) + 1  # +1 for header
        max_col = len(cols_ordered)
        stocks_sheet.auto_filter.ref = f"A1:{get_column_letter(max_col)}{max_row}"

    # Format header: bold and light fill
    header_fill = PatternFill(start_color="D9D9D9", end_color="D9D9D9", fill_type="solid")
    header_font = Font(bold=True)
    for cell in stocks_sheet[1]:
        cell.fill = header_fill
        cell.font = header_font

    # Set column widths
    for col_idx, col_name in enumerate(cols_ordered, 1):
        # Wide columns for text
        if col_name in ["conclusion", "positives", "concerns"]:
            width = 60
        elif col_name in ["data_gaps", "stance_rules"]:
            width = 50
        else:
            # Default: adaptive
            width = min(20, max(len(str(col_name)) + 2, 12))
        stocks_sheet.column_dimensions[get_column_letter(col_idx)].width = width

    # Summary sheet: just format header
    summary_sheet = wb["Summary"]
    for cell in summary_sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
    summary_sheet.column_dimensions["A"].width = 35
    summary_sheet.column_dimensions["B"].width = 15

    wb.save(xlsx_path)
    return xlsx_path
